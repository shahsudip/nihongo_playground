import os
import sys
import re
import time
import base64
import requests
from concurrent.futures import ThreadPoolExecutor

sys.path = [p for p in sys.path if "3.11" not in p]
sys.stdout.reconfigure(encoding='utf-8')

import docx
from docx.shared import Pt
from docx.oxml.ns import qn

GEMINI_API_KEY = "AQ.Ab8RN6JW3j_WeBbJvYWrWtQYACZs7lP95A4DtglIXhhR6qAfYQ"

def call_gemini(img_path):
    with open(img_path, "rb") as image_file:
        b64_img = base64.b64encode(image_file.read()).decode('utf-8')
        
    url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.5-flash:generateContent?key={GEMINI_API_KEY}"
    
    prompt = """You are a Japanese language exam image transcriber. Your task is to view this JLPT N3 Moji/Goi or Grammar question image and transcribe it into structured markdown.
FORMAT RULES:
- Ignore any header like "第X回 模擬テスト".
- Moji/Goi section: `### 文字・語彙 (Moji, Goi)` (only if you see it on the page)
- Grammar section: `### 文法 (Grammar)` (only if you see it on the page)
- Question type headers: `#### 問題 N`
- Instructions: keep the exact instruction text after each 問題 header
- Questions: `[N] question text` (Replace boxed numbers like ①, [1], etc with [N])
- Options on separate lines exactly: `1 option1`, `2 option2`, `3 option3`, `4 option4`
- Preserve all Japanese characters and formatting.
- **IMPORTANT: For any underlined text in the image, you MUST wrap it in `<u>` and `</u>` tags. For example, if 'りんご' is underlined, write `<u>りんご</u>`.**
- For 問題2 (star questions), use: `_______ ★ _______` format
- For 問題3 (reading comprehension), transcribe the full passage.
Just output the markdown text, nothing else."""

    payload = {
        "contents": [{
            "parts": [
                {"text": prompt},
                {"inline_data": {"mime_type": "image/jpeg", "data": b64_img}}
            ]
        }],
        "generationConfig": {
            "temperature": 0.0
        }
    }
    
    for attempt in range(10): # retry more times for 503s
        try:
            resp = requests.post(url, headers={"Content-Type": "application/json"}, json=payload)
            if resp.status_code == 200:
                text = resp.json()["candidates"][0]["content"]["parts"][0].get("text", "")
                text = text.replace("```markdown", "").replace("```", "").strip()
                return text
            elif resp.status_code in [429, 503, 500]:
                time.sleep(10)
            else:
                print(f"Error {resp.status_code}: {resp.text}")
                time.sleep(5)
        except Exception as e:
            time.sleep(5)
            
    return "<!-- API ERROR -->"

def process_image(args):
    idx, fname, img_path = args
    print(f"  Requesting {fname} (Page {idx+1}/10)...")
    text = call_gemini(img_path)
    print(f"  Finished {fname} (Page {idx+1}/10).")
    return idx, fname, text

def compile_docx(md_path, docx_path):
    doc = docx.Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    first_set = True
    for line in lines:
        s = line.strip()
        if not s:
            continue

        if s.startswith('## Set'):
            if not first_set:
                doc.add_page_break()
            first_set = False
            p = doc.add_heading(level=1)
            run = p.add_run(s.lstrip('#').strip())
            run.font.name = 'MS Mincho'
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'MS Mincho')
            run.font.size = Pt(16)
            run.bold = True
        elif s.startswith('####'):
            p = doc.add_heading(level=3)
            run = p.add_run(s.lstrip('#').strip())
            run.font.name = 'MS Mincho'
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'MS Mincho')
        elif s.startswith('###'):
            p = doc.add_heading(level=2)
            run = p.add_run(s.lstrip('#').strip())
            run.font.name = 'MS Mincho'
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'MS Mincho')
        elif s.startswith('#'):
            p = doc.add_heading(level=1)
            run = p.add_run(s.lstrip('#').strip())
        else:
            p = doc.add_paragraph()
            # simple parser to handle <u>...</u>
            parts = re.split(r'(<u>.*?</u>)', s)
            for part in parts:
                if part.startswith('<u>') and part.endswith('</u>'):
                    run = p.add_run(part[3:-4])
                    run.underline = True
                else:
                    run = p.add_run(part)
                run.font.name = 'MS Mincho'
                run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'MS Mincho')
                run.font.size = Pt(11)

    doc.save(docx_path)

def main():
    ws = os.path.dirname(os.path.abspath(__file__))
    img_dir = os.path.join(ws, '15setquestion')
    md_path = os.path.join(ws, '15_sets_question.md')
    docx_path = os.path.abspath(os.path.join(ws, '..', '15 sets question.docx'))

    imgs = []
    for f in os.listdir(img_dir):
        if f.lower().endswith('.jpg'):
            m = re.search(r'(\d+)', f)
            imgs.append((int(m.group(1)) if m else 9999, f))
    imgs.sort()
    
    sets = {i + 1: imgs[i * 10:(i + 1) * 10] for i in range(15)}

    # Rewrite entirely to add underlines for everything (Sets 1-15)
    with open(md_path, 'w', encoding='utf-8') as md:
        md.write("# Nihongo Power Drill N3 Moji Goi - 15 Sets\n\n")

    with open(md_path, 'a', encoding='utf-8') as md:
        for sn in range(1, 16):
            si = sets.get(sn, [])
            if not si:
                continue
            print(f"\n--- Set {sn} ---")
            md.write(f"\n## Set {sn}\n\n")

            tasks = []
            for idx, (num, fname) in enumerate(si):
                path = os.path.join(img_dir, fname)
                tasks.append((idx, fname, path))
            
            results = []
            with ThreadPoolExecutor(max_workers=3) as executor:
                for res in executor.map(process_image, tasks):
                    results.append(res)
                    
            results.sort(key=lambda x: x[0])
            for idx, fname, text in results:
                md.write(text + "\n\n")
            md.flush()
            time.sleep(2)

    print(f"\nMarkdown complete: {md_path}")
    compile_docx(md_path, docx_path)
    print(f"DOCX complete: {docx_path}")

if __name__ == "__main__":
    main()
